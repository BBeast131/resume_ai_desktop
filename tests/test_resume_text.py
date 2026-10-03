"""Reading the original resume and the prompt file."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfWriter

from app.services import resume_text
from app.services.resume_text import (
    MAX_PROMPT_CHARS,
    MAX_RESUME_CHARS,
    ResumeTextError,
    clean_extracted_text,
    extract_resume_text,
    read_prompt_file,
)


def write_pdf(path: Path, pages: list[list[str]]) -> None:
    """A minimal text PDF, written by hand: one Helvetica line per entry."""
    objects: list[bytes] = []
    kids = " ".join(f"{4 + 2 * index} 0 R" for index in range(len(pages)))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    for index, lines in enumerate(pages):
        parts = ["BT /F1 12 Tf 72 720 Td 14 TL"]
        for line in lines:
            escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            parts.append(f"({escaped}) Tj T*")
        parts.append("ET")
        stream = "\n".join(parts).encode("cp1252")
        content_number = 5 + 2 * index
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_number} 0 R"
            " /Resources << /Font << /F1 3 0 R >> >> >>".encode()
        )
        objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")

    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))


# -- text files ---------------------------------------------------------------


def test_reads_a_utf8_text_resume(tmp_path: Path):
    path = tmp_path / "resume.txt"
    path.write_text("Zoë Doe\nEngineer — 10 yrs\n", encoding="utf-8")
    assert extract_resume_text(path) == "Zoë Doe\nEngineer — 10 yrs"


def test_strips_a_utf8_bom(tmp_path: Path):
    path = tmp_path / "resume.md"
    path.write_bytes(b"\xef\xbb\xbf# Jane Doe\r\n\r\nEngineer\r\n")
    assert extract_resume_text(path) == "# Jane Doe\n\nEngineer"


def test_reads_utf16_with_a_bom(tmp_path: Path):
    path = tmp_path / "resume.txt"
    path.write_bytes("Jane Doe\nEngineer".encode("utf-16"))
    assert extract_resume_text(path) == "Jane Doe\nEngineer"


def test_falls_back_to_windows_1252(tmp_path: Path):
    path = tmp_path / "resume.TXT"
    path.write_bytes("Zoë Doe — “Engineer”".encode("cp1252"))
    assert extract_resume_text(path) == "Zoë Doe — “Engineer”"


def test_normalises_like_the_reference():
    raw = "  Jane   Doe \t\r\n\r\n\r\n\r\n   Engineer\t\tat  Acme  \n- Python \n\n"
    assert clean_extracted_text(raw) == "Jane Doe\n\nEngineer at Acme\n- Python"


def test_empty_text_file_is_refused(tmp_path: Path):
    path = tmp_path / "resume.txt"
    path.write_text(" \n\t\n", encoding="utf-8")
    with pytest.raises(ResumeTextError, match="empty"):
        extract_resume_text(path)


def test_oversized_text_is_refused_with_the_counts(tmp_path: Path):
    path = tmp_path / "resume.txt"
    path.write_text("a" * (MAX_RESUME_CHARS + 1), encoding="utf-8")
    with pytest.raises(ResumeTextError) as caught:
        extract_resume_text(path)
    assert str(caught.value) == "Extracted text is 250,001 characters. The limit is 250,000."

    path.write_text("a" * MAX_RESUME_CHARS, encoding="utf-8")
    assert len(extract_resume_text(path)) == MAX_RESUME_CHARS == 250_000


def test_a_file_over_8_mb_is_refused_before_it_is_read(tmp_path: Path):
    path = tmp_path / "resume.pdf"
    path.write_bytes(b"%PDF-1.4\n" + b"0" * (8 * 1024 * 1024))
    with pytest.raises(ResumeTextError, match="too large"):
        extract_resume_text(path)


def test_binary_data_in_a_text_file_is_refused(tmp_path: Path):
    path = tmp_path / "resume.txt"
    path.write_bytes(b"PK\x03\x04\x00\x00\x00\x00binary")
    with pytest.raises(ResumeTextError, match="text file"):
        extract_resume_text(path)


def test_unsupported_and_missing_files(tmp_path: Path):
    other = tmp_path / "resume.rtf"
    other.write_text("x", encoding="utf-8")
    with pytest.raises(ResumeTextError, match="Choose a"):
        extract_resume_text(other)
    with pytest.raises(ResumeTextError, match="could not be found"):
        extract_resume_text(tmp_path / "missing.pdf")


# -- docx ---------------------------------------------------------------------


def test_reads_a_docx_with_tables_in_reading_order(tmp_path: Path):
    document = Document()
    document.add_heading("Jane Doe", level=1)
    document.add_paragraph("Senior Engineer\tAustin, TX")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Acme Corp"
    table.cell(0, 1).text = "2019 – 2024"
    table.cell(1, 0).text = "Built retrieval pipelines"
    table.cell(1, 1).merge(table.cell(1, 0))
    nested = table.cell(0, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Remote"
    paragraph = document.add_paragraph("Python")
    paragraph.add_run().add_break()
    paragraph.add_run("SQL")
    path = tmp_path / "resume.docx"
    document.save(str(path))

    text = extract_resume_text(path)
    lines = text.split("\n")
    assert lines[0] == "Jane Doe"
    assert "Senior Engineer\tAustin, TX" in lines
    for expected in ("Acme Corp", "2019 – 2024", "Remote", "Built retrieval pipelines", "Python", "SQL"):
        assert expected in lines
    assert lines.index("Acme Corp") < lines.index("2019 – 2024") < lines.index("Built retrieval pipelines")
    assert lines.index("Built retrieval pipelines") < lines.index("Python") < lines.index("SQL")
    assert text.count("Built retrieval pipelines") == 1, "a merged cell is read once"


def test_a_vertically_merged_cell_is_read_once(tmp_path: Path):
    document = Document()
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Acme Corp"
    table.cell(0, 0).merge(table.cell(2, 0))
    for index, text in enumerate(("Engineer", "Senior Engineer", "Staff Engineer")):
        table.cell(index, 1).text = text
    path = tmp_path / "resume.docx"
    document.save(str(path))
    assert extract_resume_text(path) == "Acme Corp\nEngineer\nSenior Engineer\nStaff Engineer"


def test_an_empty_docx_is_refused(tmp_path: Path):
    path = tmp_path / "resume.docx"
    Document().save(str(path))
    with pytest.raises(ResumeTextError, match="Could not read any text"):
        extract_resume_text(path)


def test_a_file_that_is_not_really_a_docx_is_refused(tmp_path: Path):
    path = tmp_path / "resume.docx"
    path.write_bytes(b"this is not a zip archive")
    with pytest.raises(ResumeTextError, match="valid .docx"):
        extract_resume_text(path)


# -- pdf ----------------------------------------------------------------------


def test_reads_a_pdf_page_by_page(tmp_path: Path):
    path = tmp_path / "resume.pdf"
    write_pdf(path, [["Jane Doe", "Senior Engineer (AI)", "jane@example.com"], ["Experience", "Acme Corp 2019-2024"]])
    text = extract_resume_text(path)
    assert text == "Jane Doe\nSenior Engineer (AI)\njane@example.com\n\nExperience\nAcme Corp 2019-2024"


def test_a_pdf_without_text_is_refused(tmp_path: Path):
    path = tmp_path / "scan.pdf"
    write_pdf(path, [[]])
    with pytest.raises(ResumeTextError, match="Scanned image PDFs are not supported"):
        extract_resume_text(path)


def test_a_broken_pdf_is_refused_in_plain_words(tmp_path: Path):
    path = tmp_path / "resume.pdf"
    path.write_bytes(b"not a pdf at all")
    with pytest.raises(ResumeTextError) as caught:
        extract_resume_text(path)
    assert str(caught.value) == "Could not read that PDF."


def test_a_password_protected_pdf_is_refused(tmp_path: Path):
    plain = tmp_path / "plain.pdf"
    write_pdf(plain, [["Jane Doe"]])
    locked = tmp_path / "resume.pdf"
    writer = PdfWriter(clone_from=str(plain))
    writer.encrypt(user_password="secret", owner_password="owner", algorithm="RC4-128")
    writer.write(str(locked))
    with pytest.raises(ResumeTextError, match="password-protected"):
        extract_resume_text(locked)

    # Protected against editing only (no password to open): still readable.
    writer = PdfWriter(clone_from=str(plain))
    writer.encrypt(user_password="", owner_password="owner", algorithm="RC4-128")
    writer.write(str(locked))
    assert extract_resume_text(locked) == "Jane Doe"


def test_error_messages_never_quote_the_file(tmp_path: Path):
    path = tmp_path / "resume.txt"
    path.write_text("SECRET-RESUME-LINE " * 20_000, encoding="utf-8")
    with pytest.raises(ResumeTextError) as caught:
        extract_resume_text(path)
    assert "SECRET" not in str(caught.value)


# -- prompt file --------------------------------------------------------------


def test_reads_a_prompt_file(tmp_path: Path):
    path = tmp_path / "Anthony-resume-prompt.txt"
    path.write_bytes(
        b"\xef\xbb\xbf\r\nYou are a resume engine.\r\n\r\n{{PASTE_ORIGINAL_RESUME_HERE}}\r\n  keep   spacing\r\n\r\n"
    )
    assert read_prompt_file(path) == "You are a resume engine.\n\n{{PASTE_ORIGINAL_RESUME_HERE}}\n  keep   spacing"


def test_prompt_file_limits(tmp_path: Path):
    path = tmp_path / "prompt.md"
    path.write_text("p" * MAX_PROMPT_CHARS, encoding="utf-8")
    assert len(read_prompt_file(path)) == MAX_PROMPT_CHARS == 200_000

    path.write_text("p" * (MAX_PROMPT_CHARS + 1), encoding="utf-8")
    with pytest.raises(ResumeTextError) as caught:
        read_prompt_file(path)
    assert str(caught.value) == "The prompt is 200,001 characters. The limit is 200,000."

    path.write_text("\n \n", encoding="utf-8")
    with pytest.raises(ResumeTextError, match="empty"):
        read_prompt_file(path)


def test_prompt_file_must_be_text(tmp_path: Path):
    path = tmp_path / "prompt.pdf"
    path.write_bytes(b"%PDF-1.4")
    with pytest.raises(ResumeTextError, match=r"\.txt or \.md"):
        read_prompt_file(path)
    with pytest.raises(ResumeTextError, match="could not be found"):
        read_prompt_file(tmp_path / "missing.txt")


def test_module_limits_match_the_reference():
    assert resume_text.MAX_RESUME_BYTES == 8 * 1024 * 1024
    assert resume_text.MAX_RESUME_CHARS == 250_000
